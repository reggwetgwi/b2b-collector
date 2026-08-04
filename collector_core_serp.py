"""
Универсальный сборщик базы бизнесов через SERP-сервисы
"""

import os
import json
import time
import re
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Устанавливаю openpyxl...")
    os.system("pip install openpyxl")
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

try:
    import requests
except ImportError:
    print("Устанавливаю requests...")
    os.system("pip install requests")
    import requests

from serp_rotator import SERPRotator
from gemini_enricher import GeminiEnricher
from dadata_enricher import DaDataEnricher
from twogis_enricher import TwoGISEnricher
from gemini_filter import GeminiFilter


def extract_business_from_serp(result: dict, position: int) -> dict:
    """Извлечь данные бизнеса из SERP результата"""
    title = result.get("title", "")
    link = result.get("link", "")
    snippet = result.get("snippet", "")

    # Фильтруем мусорные фразы в title (рейтинги, подборки, обзоры)
    spam_phrases = [
        "рейтинг", "лучших", "лучшие", "лучшая", "топ", "топ-",
        "выбрать", "как выбрать", "сравнение", "обзор", "отзывы о",
        "список", "каталог", "справочник", "подборка",
        "стоматологи города", "стоматологи москвы",
        "стоматология в москве", "стоматология в новосибирске",
        "стоматология в санкт", "стоматология в екатеринбурге",
        "сеть стоматологических клиник"
    ]

    title_lower = title.lower()

    # Проверяем точное совпадение для общих фраз
    if title_lower in spam_phrases:
        return None

    # Проверяем вхождение для остальных
    for phrase in spam_phrases:
        if len(phrase) < 20 and phrase in title_lower:  # Короткие фразы
            return None

    link_lower = link.lower()

    # Извлекаем домен и нормализуем ссылку до главной страницы (homepage)
    domain = ""
    website = link
    if link:
        try:
            from urllib.parse import urlparse
            parsed = urlparse(link)
            domain = parsed.netloc.replace("www.", "").lower()
            if parsed.scheme and parsed.netloc:
                website = f"{parsed.scheme}://{parsed.netloc}/"
        except Exception:
            pass

    # Чёрный список доменов-агрегаторов/каталогов/справочников.
    # Сравниваем по домену (точное совпадение или поддомен), а не по подстроке,
    # чтобы не зарубить, например, "stomos.ru" из-за "mos.ru".
    blocked_domains = [
        "2gis.ru", "2gis.com", "go.2gis.com",
        "mos.ru", "gosuslugi.ru",
        "6750000.ru",
        "prodoctorov.ru", "napopravku.ru", "docdoc.ru", "sberhealth.ru",
        "zoon.ru", "yell.ru", "spr.ru", "blizko.ru", "orgpage.ru",
        "irecommend.ru", "otzovik.com", "flamp.ru", "otzyvru.com",
        "avito.ru", "youla.ru", "yandex.ru", "ya.ru", "google.com",
        "vk.com", "ok.ru", "t.me", "instagram.com", "facebook.com",
        "wikipedia.org", "rusprofile.ru", "list-org.com", "zachestnyibiznes.ru",
        "dzen.ru", "vc.ru", "rbc.ru", "kp.ru",
    ]
    for bd in blocked_domains:
        if domain == bd or domain.endswith("." + bd):
            return None

    # Чёрный список по фрагменту пути (статьи, новости, страницы поиска)
    blocked_path_markers = [
        "/articles/", "/article/", "/blog/", "/news/", "/stati/",
        "/search", "/poisk", "/rating", "/reyting", "/otzyv",
    ]
    for marker in blocked_path_markers:
        if marker in link_lower:
            return None

    # Очищаем название
    clean_name = title

    # Убираем города в конце
    cities = [" в Москве", " в СПб", " в Санкт-Петербурге", " в Екатеринбурге",
              " в Новосибирске", " в Казани", " в Красноярске", " в Нижнем Новгороде"]
    for city in cities:
        if clean_name.endswith(city):
            clean_name = clean_name[:-len(city)]

    # Убираем всё после разделителей
    separators = [" - ", " — ", " – ", " | ", ": ", "•", " — официальный", " | Платная"]
    for sep in separators:
        if sep in clean_name:
            parts = clean_name.split(sep)
            if parts[0].strip() and len(parts[0].strip()) > 3:
                clean_name = parts[0].strip()
                break

    # Убираем кавычки и лишние символы
    clean_name = clean_name.strip("«»\"'.,")

    # Если название слишком длинное (>70 символов) — обрезаем
    if len(clean_name) > 70:
        clean_name = clean_name[:70].strip()

    # Извлекаем телефон
    phone = ""
    phone_patterns = [
        r'\+7[\s\-]?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}',
        r'8[\s\-]?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}',
    ]
    for pattern in phone_patterns:
        match = re.search(pattern, snippet)
        if match:
            phone = match.group(0)
            break

    return {
        "name": clean_name,
        "original_title": title,
        "domain": domain,
        "position": position,
        "website": website,
        "phone": phone,
        "snippet": snippet,
    }


def collect_city(rotator: SERPRotator, business_type: str, city_name: str, limit: int,
                 gemini_filter=None, dadata_enricher=None, twogis_enricher=None, gemini_enricher=None,
                 log_callback=None, progress_callback=None):
    """Собрать данные по городу через SERP-сервисы"""

    if log_callback:
        log_callback(f"{'='*50}")
        log_callback(f"{city_name} — собираю {limit} бизнесов...")
        log_callback(f"{'='*50}")

    query = f"{business_type} {city_name}"
    all_businesses = []

    try:
        # Поиск в Google
        if log_callback:
            log_callback(f"🔍 Запрос в Google: {query}")

        google_data = rotator.search_google(query, log_callback=log_callback)
        google_results = google_data.get("results", [])

        if log_callback:
            log_callback(f"✓ Google вернул {len(google_results)} результатов")

        # Фильтруем через Gemini (если включено)
        if gemini_filter and google_results:
            if log_callback:
                log_callback(f"\n🤖 Gemini фильтрация Google результатов...")
            google_results = gemini_filter.filter_serp_results(google_results, business_type, city_name, log_callback=log_callback)

        # Поиск в Яндекс
        if log_callback:
            log_callback(f"🔍 Запрос в Яндекс: {query}")

        yandex_data = rotator.search_yandex(query, limit=limit, log_callback=log_callback)
        yandex_results = yandex_data.get("results", [])

        if log_callback:
            log_callback(f"✓ Яндекс вернул {len(yandex_results)} результатов")

        # Фильтруем через Gemini (если включено)
        if gemini_filter and yandex_results:
            if log_callback:
                log_callback(f"\n🤖 Gemini фильтрация Яндекс результатов...")
            yandex_results = gemini_filter.filter_serp_results(yandex_results, business_type, city_name, log_callback=log_callback)

        # Объединяем результаты (по Google, дополняем Яндекс позициями).
        # Ключ дедупликации — домен компании (а не полный URL),
        # чтобы один бизнес не задваивался из-за разных страниц одного сайта.
        businesses_map = {}

        def _init_business(b, city):
            b["yandex_rank"] = b.get("yandex_rank", "")
            b["google_rank"] = b.get("google_rank", "")
            b["city"] = city
            b["address"] = ""
            b["inn"] = ""
            b["director_phone"] = ""
            b["director_name"] = ""
            b["yandex_rating"] = ""
            b["twogis_rating"] = ""
            b["zoon_rating"] = ""
            b["services"] = ""
            return b

        for g_result in google_results[:limit]:
            business = extract_business_from_serp(g_result, g_result["position"])

            # Пропускаем агрегаторы
            if business is None:
                continue

            key = business["domain"] or business["website"] or business["name"]
            if key in businesses_map:
                continue

            business["google_rank"] = g_result["position"]
            _init_business(business, city_name)
            businesses_map[key] = business

        # Дополняем позициями из Яндекс
        for y_result in yandex_results:
            business = extract_business_from_serp(y_result, y_result["position"])

            # Пропускаем агрегаторы
            if business is None:
                continue

            key = business["domain"] or business["website"] or business["name"]
            if key in businesses_map:
                # Тот же бизнес уже есть из Google — проставляем позицию Яндекса
                if not businesses_map[key].get("yandex_rank"):
                    businesses_map[key]["yandex_rank"] = y_result["position"]
            else:
                business["yandex_rank"] = y_result["position"]
                _init_business(business, city_name)
                businesses_map[key] = business

        all_businesses = list(businesses_map.values())[:limit]

        # Добавляем ранк (порядковый номер)
        for idx, business in enumerate(all_businesses, 1):
            business["rank"] = idx

        # Обогащение через DaData (если включено)
        if dadata_enricher and all_businesses:
            if log_callback:
                log_callback(f"\n🏢 Запуск DaData обогащения для {len(all_businesses)} бизнесов...")

            all_businesses = dadata_enricher.enrich_batch(all_businesses, log_callback=log_callback)

            if log_callback:
                log_callback(f"✓ DaData обогащение завершено\n")

        # Обогащение через 2GIS (если включено)
        if twogis_enricher and all_businesses:
            if log_callback:
                log_callback(f"\n📍 Запуск 2GIS обогащения для {len(all_businesses)} бизнесов...")

            all_businesses = twogis_enricher.enrich_batch(all_businesses, log_callback=log_callback)

            if log_callback:
                log_callback(f"✓ 2GIS обогащение завершено\n")

        # Обогащение через Gemini (если включено и предыдущие не заполнили всё)
        if gemini_enricher and all_businesses:
            if log_callback:
                log_callback(f"\n🤖 Запуск Gemini обогащения для {len(all_businesses)} бизнесов...")

            all_businesses = gemini_enricher.enrich_batch(all_businesses, log_callback=log_callback)

            if log_callback:
                log_callback(f"✓ Gemini обогащение завершено\n")

        if progress_callback:
            progress_callback(len(all_businesses), limit)

        if log_callback:
            log_callback(f"✓ Итого для {city_name}: {len(all_businesses)} бизнесов")

    except Exception as e:
        if log_callback:
            log_callback(f"❌ ОШИБКА: {e}")

    return all_businesses


def save_xlsx(all_businesses, business_type, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Бизнесы"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    city_fills = {}
    city_colors = [
        "D6E4F0", "D5F5E3", "FEF9E7", "FDEDEC", "EBF5FB",
        "F5EEF8", "FDF2E9", "E8F8F5", "FDFEFE", "EAF2FF",
    ]

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["№", "Название", "Город", "Адрес", "Телефон", "ИНН", "Директор (ФИО)", "Тел. директора", "Google↑", "Яндекс↑", "Яндекс (0-5)", "2ГИС (0-5)", "Zoon (0-10)", "Сайт", "Услуги"]
    col_widths = [5, 35, 18, 40, 20, 15, 30, 20, 10, 10, 14, 13, 13, 30, 45]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30

    color_idx = 0
    for row_idx, business in enumerate(all_businesses, 2):
        city = business.get("city", "")
        if city not in city_fills:
            color = city_colors[color_idx % len(city_colors)]
            city_fills[city] = PatternFill("solid", start_color=color)
            color_idx += 1
        fill = city_fills[city]

        row_data = [
            business.get("rank", row_idx - 1),
            business.get("name", ""),
            city,
            business.get("address", ""),
            business.get("phone", ""),
            business.get("inn", ""),
            business.get("director_name", ""),
            business.get("director_phone", ""),
            business.get("google_rank", ""),
            business.get("yandex_rank", ""),
            business.get("yandex_rating", ""),
            business.get("twogis_rating", ""),
            business.get("zoon_rating", ""),
            business.get("website", ""),
            business.get("services", ""),
        ]

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.border = border
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx in [4, 15]))

            if col_idx in [1, 9, 10, 11, 12, 13]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            if col_idx == 14 and value:
                url = value if value.startswith("http") else f"https://{value}"
                cell.hyperlink = url
                cell.font = Font(name="Arial", size=10, color="1F4E79", underline="single")

        ws.row_dimensions[row_idx].height = 40

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{len(all_businesses) + 1}"

    ws_stat = wb.create_sheet("Статистика")
    ws_stat["A1"] = "Сводка по городам"
    ws_stat["A1"].font = Font(name="Arial", bold=True, size=13, color="1F4E79")

    stat_headers = ["Город", "Кол-во бизнесов", "Средний рейтинг"]
    for col_idx, h in enumerate(stat_headers, 1):
        cell = ws_stat.cell(row=2, column=col_idx, value=h)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color="1F4E79")
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    city_stats = {}
    for business in all_businesses:
        city = business.get("city", "Неизвестно")
        if city not in city_stats:
            city_stats[city] = {"count": 0, "ratings": []}
        city_stats[city]["count"] += 1
        try:
            r = float(str(business.get("yandex_rating", "")).replace(",", "."))
            city_stats[city]["ratings"].append(r)
        except (ValueError, TypeError):
            pass

    for stat_row, (city, stats) in enumerate(city_stats.items(), 3):
        avg_rating = (
            round(sum(stats["ratings"]) / len(stats["ratings"]), 2)
            if stats["ratings"] else "—"
        )
        fill = city_fills.get(city, PatternFill("solid", start_color="FFFFFF"))
        for col_idx, value in enumerate([city, stats["count"], avg_rating], 1):
            cell = ws_stat.cell(row=stat_row, column=col_idx, value=value)
            cell.fill = fill
            cell.border = border
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")

    ws_stat.column_dimensions["A"].width = 22
    ws_stat.column_dimensions["B"].width = 18
    ws_stat.column_dimensions["C"].width = 20

    total_row = len(city_stats) + 3
    ws_stat.cell(row=total_row, column=1, value="ИТОГО").font = Font(name="Arial", bold=True)
    ws_stat.cell(row=total_row, column=2, value=len(all_businesses)).font = Font(name="Arial", bold=True)
    for col_idx in range(1, 4):
        ws_stat.cell(row=total_row, column=col_idx).border = border
        ws_stat.cell(row=total_row, column=col_idx).fill = PatternFill("solid", start_color="D5D8DC")

    wb.save(filename)
    return filename
