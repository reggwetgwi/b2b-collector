"""
ДВИЖОК сборщика базы бизнесов (release).
Конфиг-driven, с колбэками для GUI (log, progress). Самодостаточный.

Слои (включаются флагами cfg):
  • registry  — скачать архив реестра МСП ФНС и отфильтровать по ОКВЭД+регион → 100% ИНН
  • web       — топ по популярности из Яндекса (XMLRiver), сайт→бренд/тел/email/мессенджеры/ИНН/ОГРН
  • playwright — JS-рендер сайтов ради ИНН из JS-футеров
  • dadata    — findById по ИНН → директор/адрес/статус
  • twogis    — рейтинг/отзывы/адрес 2ГИС

Итог — Excel со всеми полями.
"""
import os, re, json, time, zipfile, html as _htmllib, warnings
# Playwright ищет встроенный chromium внутри пакета (портативно для PyInstaller)
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "0")
import requests
import xml.etree.ElementTree as ET
from urllib.parse import urljoin, urlparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

MSP_URL = "https://file.nalog.ru/opendata/7707329152-rsmp/data-10052026-structure-10062025.zip"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"}

# регионы Яндекса (lr) и КЛАДР/коды для основных городов
CITY_DB = {
    "москва":           {"lr": 213, "kladr": "7700000000000", "regions": ["77", "50"], "guard": ["москва", "московск"]},
    "санкт-петербург":  {"lr": 2,   "kladr": "7800000000000", "regions": ["78", "47"], "guard": ["санкт-петербург", "ленинградск"]},
    "новосибирск":      {"lr": 65,  "kladr": "5400000100000", "regions": ["54"],       "guard": ["новосибирск"]},
    "екатеринбург":     {"lr": 54,  "kladr": "6600000100000", "regions": ["66"],       "guard": ["екатеринбург", "свердловск"]},
    "казань":           {"lr": 43,  "kladr": "1600000100000", "regions": ["16"],       "guard": ["казань", "татарстан"]},
    "нижний новгород":  {"lr": 47,  "kladr": "5200000100000", "regions": ["52"],       "guard": ["нижний новгород", "нижегородск"]},
}

BLOCKED_DOMAINS = [
    "2gis.ru", "gis.ru", "mos.ru", "gosuslugi.ru", "6750000.ru", "sravni.ru", "32top.ru",
    "prodoctorov.ru", "napopravku.ru", "docdoc.ru", "sberhealth.ru", "topdent.ru",
    "zoon.ru", "yell.ru", "spr.ru", "blizko.ru", "orgpage.ru", "firmika.ru", "doctu.ru",
    "irecommend.ru", "otzovik.com", "flamp.ru", "otzyvru.com", "gidmed.com", "medbooking.com",
    "avito.ru", "youla.ru", "yandex.ru", "ya.ru", "google.com", "infodoctor.ru", "medportal.ru",
    "vk.com", "ok.ru", "t.me", "instagram.com", "facebook.com", "stom-firms.ru", "expomap.ru",
    "wikipedia.org", "rusprofile.ru", "list-org.com", "zachestnyibiznes.ru", "doktornarabote.ru",
    "dzen.ru", "vc.ru", "rbc.ru", "kp.ru", "youtube.com",
]
BLOCKED_PATH = ["/articles/", "/article/", "/blog/", "/news/", "/search", "/rating"]
GOOD_TLD = {"ru", "com", "su", "net", "org", "moscow", "online", "info", "pro", "bio", "me",
            "clinic", "dental", "site", "store", "shop", "rus", "spb", "su"}
_STOP = {"медицина","медицинский","медицинская","стоматология","стоматологическая","стоматологический",
         "клиника","клиник","клиники","центр","дентал","дент","стом","смайл","smile","dental","clinic",
         "ооо","ао","зао","фгбу","гауз","фирма","группа","групп","сеть"}


# ───────────────────── чистые утилиты ─────────────────────
def domain_of(url):
    try: return urlparse(url).netloc.replace("www.", "").lower()
    except Exception: return ""

def inn_ck(inn):
    if not inn or not inn.isdigit(): return False
    d = [int(x) for x in inn]
    if len(inn) == 10:
        c = [2,4,10,3,5,9,4,6,8]; return (sum(c[i]*d[i] for i in range(9))%11)%10 == d[9]
    if len(inn) == 12:
        c1=[7,2,4,10,3,5,9,4,6,8]; c2=[3,7,2,4,10,3,5,9,4,6,8]
        return ((sum(c1[i]*d[i] for i in range(10))%11)%10==d[10] and (sum(c2[i]*d[i] for i in range(11))%11)%10==d[11])
    return False

def ogrn_ck(o):
    # ОГРН-13: первые 12 цифр mod 11, последняя цифра остатка = контрольная
    # ОГРНИП-15: первые 14 цифр mod 13, последняя цифра остатка = контрольная
    if not o or not o.isdigit() or len(o) not in (13, 15): return False
    if len(o) == 13:
        return int(o[:12]) % 11 % 10 == int(o[12])
    return int(o[:14]) % 13 % 10 == int(o[14])

def _noscript(raw):
    raw = re.sub(r"<script\b[^>]*>.*?</script>", " ", raw, flags=re.S|re.I)
    return re.sub(r"<style\b[^>]*>.*?</style>", " ", raw, flags=re.S|re.I)

def strip_tags(raw):
    return re.sub(r"\s+", " ", _htmllib.unescape(re.sub(r"<[^>]+>", " ", _noscript(raw))).replace("\xa0", " "))

def find_inns(raw):
    out=[]; t=strip_tags(raw)
    for m in re.finditer(r"ИНН[^0-9]{0,6}([0-9][0-9 ]{8,16}[0-9])", t, re.I):
        d=re.sub(r"\D","",m.group(1))
        if len(d) in (10,12) and inn_ck(d): out.append(d)
    return out

def find_ogrns(raw):
    out=[]; t=strip_tags(raw)
    for m in re.finditer(r"ОГРН[ИП]?[^0-9]{0,6}([0-9][0-9 ]{11,18}[0-9])", t, re.I):
        d=re.sub(r"\D","",m.group(1))
        if ogrn_ck(d): out.append(d)
    return out

def find_emails(raw):
    raw=_noscript(raw)
    em=set(re.findall(r"mailto:([A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-zА-Яа-я]{2,})", raw, re.I))
    # латиница + кириллические домены (.москва/.рф/.онлайн и т.п.)
    em|=set(re.findall(r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.(?:[A-Za-z]{2,8}|[А-Яа-я]{2,7})\b", strip_tags(raw)))
    bad=("example.","sentry","wixpress","@2x",".png",".jpg",".gif",".svg",".webp",
         "domain.","your@","@example","noreply@","no-reply@","@sentry","@wix")
    out=[]
    for e in em:
        e=e.lower().strip(".")
        if any(b in e for b in bad): continue
        tld=e.rsplit(".",1)[-1]
        if tld not in GOOD_TLD and tld not in ("москва","рф","онлайн","дети","сайт"): continue
        # отсекаем ложные: префикс из одних цифр (2024@..., артикулы) и слишком короткий локал
        local=e.split("@",1)[0]
        if local.isdigit(): continue
        if len(local) < 2: continue
        out.append(e)
    return sorted(set(out))

def find_socials(raw):
    s=set()
    for m in re.findall(r'https?://(?:wa\.me|api\.whatsapp\.com)/[^\s"\'<>]+', raw): s.add(m.split("&")[0])
    for m in re.findall(r'https?://t\.me/[^\s"\'<>]+', raw): s.add(m)
    for m in re.findall(r'https?://(?:vk\.com|vk\.ru)/[^\s"\'<>]+', raw): s.add(m.split("?")[0])
    return list(s)[:4]

def _norm_phone(raw_match):
    """Нормализовать найденный телефон к +7XXXXXXXXXX. Возвращает '' если структура битая."""
    dg = re.sub(r"\D", "", raw_match)
    if len(dg) == 11 and dg[0] in "78":
        return "+7" + dg[1:]
    if len(dg) == 10:  # городской/мобильный без кода страны
        return "+7" + dg
    return ""

def find_phones(raw):
    t = strip_tags(raw)
    ph = []
    # строгие паттерны (как для РКН): код страны 8/+7, либо городской с кодом
    pats = [
        r"(?:\+7|8)\s*\(?\d{3,4}\)?\s*\d{3}[\-\s]?\d{2}[\-\s]?\d{2}",
        r"\(\d{3,4}\)\s*\d{3}[\-\s]?\d{2}[\-\s]?\d{2}",   # (495) 123-45-67 без 8
    ]
    for pat in pats:
        for m in re.findall(pat, t):
            p = _norm_phone(m)
            if p and p not in ph:
                ph.append(p)
    return ph[:5]

def brand_from(raw):
    for pat in (r'<meta[^>]+og:site_name[^>]+content=["\']([^"\']+)',
                r'<meta[^>]+name=["\']application-name["\'][^>]+content=["\']([^"\']+)'):
        m=re.search(pat, raw, re.I)
        if m and len(m.group(1).strip())>1: return m.group(1).strip()
    m=re.search(r"<title[^>]*>(.*?)</title>", raw, re.S|re.I)
    if m:
        t=re.sub(r"\s+"," ",m.group(1)).strip()
        for sep in [" — "," - "," | ",": "," – "]:
            if sep in t: return t.split(sep)[0].strip()
        return t[:60]
    return ""

def _tok(s):
    s=re.sub(r'[«»"\'`.,;:!?/\\\-]'," ",(s or "").lower())
    return {w for w in re.findall(r"[a-zа-я0-9]{3,}",s) if w not in _STOP}

# должности руководства, рядом с которыми ищем ФИО на /team /о-команде
LEAD_TITLES = re.compile(
    r"(главн\w*\s+врач|глав\.?\s*врач|директор|руководител|заведующ\w*|"
    r"учредител|владелец|основател|генеральн\w*\s+директор)", re.I)
FIO_RE = re.compile(r"[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+(?:\s+[А-ЯЁ][а-яё]+)?")

def find_leaders(raw):
    """Со страниц команды: ФИО + должность руководства (главврач/директор/учредитель).
    Берём ближайшее ФИО к должности (после, иначе до), окно узкое — чтобы не путать людей."""
    t = strip_tags(raw)
    leaders = []
    used_fio = set()
    for m in LEAD_TITLES.finditer(t):
        title = m.group(0).lower()
        # сначала ищем ФИО сразу ПОСЛЕ должности (40 симв), потом ДО (40 симв)
        after = t[m.end():m.end()+40]
        before = t[max(0, m.start()-40):m.start()]
        fm = FIO_RE.search(after) or FIO_RE.search(before)
        if fm:
            fio = fm.group(0)
            if len(fio.split()) < 2 or fio in used_fio:
                continue
            used_fio.add(fio)
            leaders.append(f"{fio} — {title}")
    return leaders[:5]


class Engine:
    def __init__(self, cfg, log=print, progress=lambda *a: None, should_stop=lambda: False):
        self.cfg = cfg
        self.log = log
        self.progress = progress
        self.should_stop = should_stop
        self.cache = {}

    # ---------- реестр ФНС ----------
    def _find_file(self, name, min_size=0):
        """Найти файл в текущей или родительских папках (до 4 уровней вверх)."""
        probes = [name]
        cur = os.path.abspath(".")
        for _ in range(4):
            cur = os.path.dirname(cur)
            probes.append(os.path.join(cur, name))
        for p in probes:
            if os.path.exists(p) and os.path.getsize(p) >= min_size:
                return os.path.abspath(p)
        return None

    def registry_list(self, region_codes):
        """100%% ИНН из реестра МСП ФНС. Находит существующий архив/кэш, качает только если нет."""
        okp = self.cfg["okved_prefix"]
        cache_name = f"msp_{okp}_{'_'.join(region_codes)}.json"
        cache_json = self._find_file(cache_name, min_size=100) or cache_name
        if os.path.exists(cache_json):
            self.log(f"Реестр из кэша: {os.path.basename(cache_json)}")
            return json.load(open(cache_json, encoding="utf-8"))

        zip_path = self._find_file("msp.zip", min_size=1_000_000_000) or "msp.zip"
        if os.path.exists(zip_path) and os.path.getsize(zip_path) >= 1_000_000_000:
            self.log(f"Архив реестра: {os.path.basename(zip_path)} ({os.path.getsize(zip_path)//1024//1024} МБ)")
        else:
            self.log("Качаю архив реестра МСП ФНС (~2.1 ГБ, один раз)...")
            self._download(MSP_URL, zip_path)
        self.log("Фильтрую реестр по ОКВЭД/региону...")
        out = self._parse_msp(zip_path, okp, set(region_codes))
        json.dump(out, open(cache_name, "w", encoding="utf-8"), ensure_ascii=False)
        return out

    def _download(self, url, dst):
        pos = os.path.getsize(dst) if os.path.exists(dst) else 0
        hdr = {"Range": f"bytes={pos}-"} if pos else {}
        with requests.get(url, headers=hdr, stream=True, timeout=60) as r:
            total = int(r.headers.get("Content-Length", 0)) + pos
            with open(dst, "ab" if pos else "wb") as f:
                done = pos; mark = pos + 100*1024*1024
                for ch in r.iter_content(1024*1024):
                    if self.should_stop(): return
                    f.write(ch); done += len(ch)
                    if done >= mark:
                        self.log(f"  загрузка {done//1024//1024} / {total//1024//1024} МБ")
                        self.progress(done, total); mark += 100*1024*1024

    @staticmethod
    def _ln(tag): return tag.split("}")[-1]

    def _parse_msp(self, zip_path, okp, regions):
        out=[]; seen=set()
        with zipfile.ZipFile(zip_path) as z:
            xmls=[n for n in z.namelist() if n.lower().endswith(".xml")]
            for i, fn in enumerate(xmls, 1):
                if self.should_stop(): break
                try:
                    with z.open(fn) as f:
                        for ev, elem in ET.iterparse(f, events=("end",)):
                            if self._ln(elem.tag)=="Документ":
                                rec=self._proc_doc(elem)
                                if (rec["inn"] and rec["region"] in regions
                                        and any(o.startswith(okp) for o in rec["okveds"])
                                        and rec["inn"] not in seen):
                                    seen.add(rec["inn"]); out.append(rec)
                                elem.clear()
                except ET.ParseError:
                    continue
                if i % 800 == 0:
                    self.log(f"  реестр: {i}/{len(xmls)} файлов, найдено {len(out)}")
        self.log(f"Реестр: отобрано {len(out)} записей")
        return out

    def _proc_doc(self, doc):
        flat={}; ok=set(); fam=im=ot=""; ip=False
        for n in doc.iter():
            nm=self._ln(n.tag); a=n.attrib
            if nm=="ИПВклМСП": ip=True
            if nm in ("СвОКВЭДОсн","СвОКВЭДДоп") and a.get("КодОКВЭД"): ok.add(a["КодОКВЭД"])
            if nm=="КодОКВЭД" and n.text: ok.add(n.text.strip())
            if nm=="ФИОИП": fam=a.get("Фамилия",""); im=a.get("Имя",""); ot=a.get("Отчество","")
            if n.text and n.text.strip(): flat.setdefault(nm, n.text.strip())
            for k,v in a.items(): flat.setdefault(self._ln(k), v)
        inn=flat.get("ИННЮЛ") or flat.get("ИННФЛ") or ""
        return {"inn":inn, "name":flat.get("НаимОрг") or flat.get("НаимОргСокр") or "",
                "region":flat.get("КодРегион") or "", "okveds":sorted(ok),
                "is_ip":ip, "fio":" ".join(x for x in (fam,im,ot) if x) if ip else ""}

    # ---------- web ----------
    # Временные коды ошибок xmlriver — нужен перезапрос с теми же параметрами.
    _XR_RETRY_CODES = {"20", "21", "22", "23", "24", "110", "111", "115",
                       "201", "202", "203", "500"}

    # Два источника выдачи Яндекса у xmlriver (параметры одинаковые, отличается путь).
    # Порядок = приоритет фолбэка. Основной — Search API (быстрее/дешевле, работает,
    # когда живая выдача перегружена 111); запасной — живая выдача (полная, как у пользователя).
    _YX_SOURCES = [("Search API",   "yandex/xml"),
                   ("живая выдача", "search_yandex/xml")]

    def yandex_page(self, query, page, lr):
        # Порядок источников с «залипанием» на рабочем: один раз поймав мёртвый
        # эндпоинт, дальше по умолчанию ходим на запасной (а если и он умрёт — обратно).
        order = getattr(self, "_yx_order", None)
        if order is None:
            order = list(range(len(self._YX_SOURCES)))
            self._yx_order = order

        last_code = None
        for pos, src_i in enumerate(list(order)):
            label, path = self._YX_SOURCES[src_i]
            attempts = 3 if pos == 0 else 2  # основной пытаем дольше, запасной — пара попыток
            res, code, ok = self._yx_request(path, query, page, lr, attempts, label)
            if ok:
                if src_i != self._yx_order[0]:
                    self._yx_order.remove(src_i)
                    self._yx_order.insert(0, src_i)
                    self.log(f"  ↪ Яндекс: переключаюсь на источник «{label}»")
                return res
            last_code = code
            if pos == 0 and len(order) > 1:
                self.log(f"  ⚠️ Яндекс: «{label}» недоступна (код {code}), пробую запасной источник…")
        self.log(f"  ⚠️ Яндекс: страница {page} не получена ни одним источником (последний код {last_code})")
        return []

    def _yx_request(self, path, query, page, lr, attempts, label):
        """Запрос одной страницы к одному источнику.

        Возвращает (results, last_code, ok). ok=False — источник недоступен,
        нужно пробовать другой. ok=True с пустым списком — валидный «нет результатов».
        """
        u=(f"https://xmlriver.com/{path}?user={self.cfg['xr_user']}&key={self.cfg['xr_key']}"
           f"&query={requests.utils.quote(query)}&groupby=10&lr={lr}&domain=ru&page={page}")
        last_code=None
        for attempt in range(attempts):
            try:
                r=requests.get(u, timeout=60)
                m=re.search(r'<error code="(\d+)"', r.text)
                if m:
                    code=m.group(1); last_code=code
                    if code=="15":
                        return [], code, True   # нет результатов — валидный пустой ответ
                    if code=="200":
                        self.log(f"  ⛔ Яндекс/{label}: на счету xmlriver кончились деньги (код 200)")
                        return [], code, False
                    if code in self._XR_RETRY_CODES:
                        # временная ошибка — ждём и повторяем тем же запросом
                        time.sleep(min(20, 4*(2**attempt)))
                        continue
                    # постоянная ошибка этого источника (неверный параметр/ключ/IP) — повтор не поможет
                    return [], code, False
                root=ET.fromstring(r.content); res=[]
                for idx,g in enumerate(root.findall(".//group")):
                    d=g.find(".//doc")
                    if d is None: continue
                    uu=d.find(".//url"); tt=d.find(".//title")
                    res.append({"url":uu.text.strip() if uu is not None and uu.text else "",
                                "title":"".join(tt.itertext()).strip() if tt is not None else "",
                                "position":page*10+idx+1})
                return res, None, True
            except (requests.exceptions.RequestException, ET.ParseError):
                time.sleep(min(20, 4*(2**attempt)))
        return [], (last_code or "timeout"), False

    def _blocked(self, url, dom, title, city):
        low=url.lower()
        if any(dom==b or dom.endswith("."+b) for b in BLOCKED_DOMAINS): return True
        if any(m in low for m in BLOCKED_PATH): return True
        tl=(title or "").lower()
        if re.search(r"(\d+\s*(лучш|клиник)|рейтинг|лучши|топ[\s\-]?\d|подборк|обзор|сравнен|каталог|где\s+сделать|как\s+выбрать)", tl): return True
        return False

    def discover(self, query, city, lr, target, max_pages):
        seen={}; rank=0; q=0
        for syn_i, seed in enumerate([query] + [f"{query} {s}" for s in self.cfg.get("synonyms", [])]):
            if len(seen)>=target or self.should_stop(): break
            empty=0
            for page in range(max_pages):
                if self.should_stop(): break
                docs=self.yandex_page(seed, page, lr); q+=1
                if not docs: break
                new=0
                for d in docs:
                    dom=domain_of(d["url"])
                    if not dom or dom in seen or self._blocked(d["url"],dom,d["title"],city): continue
                    rank+=1
                    seen[dom]={"domain":dom,"yandex_pos":d["position"],"rank":rank,"title":d["title"]}
                    new+=1
                empty=empty+1 if new==0 else 0
                if empty>=(3 if syn_i==0 else 1) or len(seen)>=target: break
                time.sleep(0.2)
            self.log(f"  Яндекс «{seed}»: всего {len(seen)}")
        self.log(f"  запросов Яндекс: {q} (~{q*0.025:.2f}₽)")
        return list(seen.values())

    def _get(self, url):
        if url in self.cache: return self.cache[url]
        out=None
        try:
            r=requests.get(url, headers=UA, timeout=12, verify=False)
            if r.status_code<400:
                r.encoding=r.apparent_encoding or r.encoding; out=r.text
        except requests.exceptions.RequestException: pass
        self.cache[url]=out; return out

    def _pdf_inns(self, url):
        try:
            import pypdf, io
            r=requests.get(url, headers=UA, timeout=15, verify=False)
            if r.status_code>=400: return [],[]
            rd=pypdf.PdfReader(io.BytesIO(r.content))
            txt=" ".join((p.extract_text() or "") for p in rd.pages[:6])
            return find_inns(txt), find_ogrns(txt)
        except Exception: return [],[]

    def enrich_site(self, b):
        dom=b["domain"]; b.update({"site_brand":"","phones":[],"emails":[],"socials":[],"inn":"","ogrn":""})
        base=home=None
        for sch in ("https://","http://"):
            home=self._get(f"{sch}{dom}/")
            if home is not None: base=f"{sch}{dom}"; break
        if home is None: return b
        b["site_brand"]=brand_from(home)
        b["leaders"]=[]
        inns=list(find_inns(home)); ogrns=list(find_ogrns(home))
        b["phones"]=find_phones(home); b["emails"]=find_emails(home); b["socials"]=find_socials(home)
        b["leaders"]=find_leaders(home)
        legal=[urljoin(base+"/",h) for h in re.findall(r'href=["\']([^"\']+)["\']',home)
               if re.search(r"(polit|privacy|konfiden|personal|rekvizit|policy|/about|/o-)",h,re.I)]
        legal+=[base+p for p in ("/contacts","/kontakty","/policy","/privacy","/politika-konfidencialnosti","/rekvizity","/about")]
        # страницы команды/руководства — маршрут 2 (ЛПР со страниц /team)
        team=[base+p for p in ("/team","/komanda","/o-klinike","/vrachi","/doctors","/specialisty","/about-us")]
        seen=set()
        for url in legal+team:
            if url in seen: continue
            seen.add(url)
            if url.lower().endswith(".pdf"):
                pi,po=self._pdf_inns(url); inns+=pi; ogrns+=po; continue
            h=self._get(url)
            if not h: continue
            inns+=find_inns(h); ogrns+=find_ogrns(h)
            if not b["emails"]: b["emails"]=find_emails(h)
            if not b["phones"]: b["phones"]=find_phones(h)
            if not b["socials"]: b["socials"]=find_socials(h)
            if not b["leaders"]: b["leaders"]=find_leaders(h)
        pref=[i for i in inns if i[:2] in self.cfg.get("region_codes", {"77","50"})]
        b["inn"]=(pref or inns)[0] if inns else ""; b["ogrn"]=(ogrns or [""])[0]
        return b

    def playwright_pass(self, businesses):
        todo=[b for b in businesses if not b.get("inn") and not b.get("ogrn")][:self.cfg.get("pw_cap",220)]
        if not todo: return
        try:
            from playwright.sync_api import sync_playwright
        except Exception as e:
            self.log(f"  Playwright не доступен ({e}) — пропускаю JS-рендер"); return
        self.log(f"  Playwright: JS-рендер {len(todo)} сайтов")
        try:
            with sync_playwright() as p:
                br=p.chromium.launch()
                for b in todo:
                    if self.should_stop(): break
                    for sch in ("https://","http://"):
                        try:
                            pg=br.new_page(user_agent=UA["User-Agent"])
                            pg.goto(f"{sch}{b['domain']}/", timeout=15000, wait_until="domcontentloaded")
                            pg.wait_for_timeout(1500); html=pg.content(); pg.close()
                        except Exception:
                            try: pg.close()
                            except Exception: pass
                            continue
                        inns=find_inns(html); ogrns=find_ogrns(html)
                        if not b["emails"]: b["emails"]=find_emails(html)
                        if inns or ogrns:
                            pref=[i for i in inns if i[:2] in self.cfg.get("region_codes", {"77","50"})]
                            b["inn"]=(pref or inns)[0] if inns else ""; b["ogrn"]=(ogrns or [""])[0]
                        break
                br.close()
        except Exception as e:
            self.log(f"  Playwright ошибка: {e}")

    # ---------- DaData / 2GIS ----------
    def dadata_find(self, ident):
        try:
            r=requests.post("https://suggestions.dadata.ru/suggestions/api/4_1/rs/findById/party",
                            headers={"Authorization":f"Token {self.cfg['dadata']}"}, json={"query":ident}, timeout=15)
            if r.status_code in (403,429): return "LIMIT"
            s=r.json().get("suggestions",[])
            if s:
                d=s[0]["data"]
                # учредители-физлица из ЕГРЮЛ (часто = реальные владельцы/ЛПР)
                founders=[]
                for f in (d.get("founders") or []):
                    nm=(f.get("fio") or {}).get("source") or f.get("name") or ""
                    if nm and "ОГРАНИЧЕННОЙ" not in nm.upper() and "ОБЩЕСТВО" not in nm.upper():
                        founders.append(nm)
                return {"name":s[0]["value"],"director":(d.get("management") or {}).get("name",""),
                        "address":(d.get("address") or {}).get("value",""),
                        "status":(d.get("state") or {}).get("status",""),
                        "founders":", ".join(founders[:3])}
        except requests.exceptions.RequestException: pass
        return None

    def twogis(self, businesses, city):
        if not self.cfg.get("twogis"): return
        try:
            from twogis_enricher import TwoGISEnricher
            en=TwoGISEnricher([self.cfg["twogis"]])
        except Exception as e:
            self.log(f"  2ГИС недоступен: {e}"); return
        def one(b):
            nm=re.sub(r'^(ООО|АО|ЗАО|ПАО)\s+',"",(b.get("egrul_name") or b.get("site_brand") or b.get("title","")).strip(' "«»'))
            pd=en.search_place(nm, city)
            if pd:
                b["twogis_rating"]=pd.get("rating",""); b["twogis_reviews"]=pd.get("reviews_count","")
                if not b.get("phones") and pd.get("phones"): b["phones"]=pd["phones"]
            return b
        with ThreadPoolExecutor(max_workers=6) as ex:
            list(as_completed([ex.submit(one,b) for b in businesses]))

    # ---------- оркестрация (веб фаза на город) ----------
    def _web_city(self, city, target):
        """Только веб-слой: discovery + сайты + Playwright. Возвращает список rows."""
        c = CITY_DB.get(city.lower(), {"lr":213,"kladr":"7700000000000","regions":["77"],"guard":[city.lower()]})
        query = f"{self.cfg['business']} {city}"
        self.log(f"=== {city}: web discovery (лимит {target}) ===")
        web = self.discover(query, city, c["lr"], target, self.cfg.get("max_pages",20))
        self.log(f"=== {city}: обогащение {len(web)} сайтов ===")
        with ThreadPoolExecutor(max_workers=12) as ex:
            futs = {ex.submit(self.enrich_site, b): b for b in web}
            for i, f in enumerate(as_completed(futs), 1):
                if i % 8 == 0 or i == len(web):
                    self.log(f"  сайты: {i}/{len(web)}")
        if self.cfg.get("playwright"): self.playwright_pass(web)
        for b in web: b["city"] = city; b["c"] = c
        return web

    # ---------- DaData чекпойнт ----------
    def _dadata_ckpt(self):
        ckpt = {}
        try:
            if os.path.exists("dadata_ckpt.json"):
                ckpt = json.load(open("dadata_ckpt.json", encoding="utf-8"))
                self.log(f"DaData чекпойнт: {len(ckpt)} записей")
        except Exception: pass
        return ckpt

    def _run_dadata(self, rows):
        """DaData findById только по контактным rows (у кого ИНН и веб-присутствие). С чекпойнтом и resume."""
        if not self.cfg.get("dadata"): return
        ckpt = self._dadata_ckpt()
        limit = False; new = 0
        # ТОЛЬКО контактные = у кого есть веб-присутствие (сайт/телефон/email).
        # Записи реестра без сайта (мёртвые лиды) в DaData НЕ шлём — иначе сожрём лимит впустую.
        def contactable(b):
            return bool(b.get("domain") or b.get("phones") or b.get("emails"))
        todo = [b for b in rows
                if contactable(b) and not b.get("director") and (b.get("inn") or b.get("ogrn"))]
        self.log(f"=== DaData: {len(todo)} контактных (из {len(rows)} всего) ===")
        for i, b in enumerate(todo, 1):
            if self.should_stop() or limit: break
            inn = b.get("inn")
            if inn in ckpt:
                b.update(ckpt[inn]); continue
            if not inn: continue
            d = self.dadata_find(inn)
            if d == "LIMIT":
                limit = True; self.log(f"  лимит — сохраняю чекпойнт ({len(ckpt)}), продолжу завтра")
                break
            if d:
                b["director"] = d["director"]; b["reg_address"] = d["address"]
                b["status"] = d["status"]; b["egrul_name"] = d.get("name") or b.get("egrul_name","")
                b["founders"] = d.get("founders", "")
                ckpt[inn] = d; new += 1
            time.sleep(0.04)
            if i % 100 == 0: self.log(f"  DaData: {i}/{len(todo)}")
        json.dump(ckpt, open("dadata_ckpt.json", "w", encoding="utf-8"), ensure_ascii=False)
        self.log(f"  DaData: +{new} (чекпойнт: {len(ckpt)})")

    # ---------- главная оркестрация ----------
    def run(self):
        all_web = []
        default_target = self.cfg.get("target", 200)
        # города могут быть с индивидуальным лимитом: {"name": "Москва", "limit": 100}
        # или строкой — тогда дефолтный лимит
        city_items = []
        for c in self.cfg["cities"]:
            if isinstance(c, dict):
                city_items.append((c["name"], int(c.get("limit") or default_target)))
            else:
                city_items.append((c, default_target))
        self.cfg["cities"] = [name for name, _ in city_items]  # для последующих фаз

        # ФАЗА 1: веб-проход по всем городам (быстро)
        for city, tgt in city_items:
            if self.should_stop(): break
            try: all_web += self._web_city(city, tgt)
            except Exception as e: self.log(f"❌ {city}: {e}")

        # ФАЗА 2: реестр один раз (все регионы скопом) + джойн с вебом
        rows = []
        if self.cfg.get("registry"):
            all_regions = set()
            for b in all_web:
                for r in (b.get("c") or {}).get("regions", []): all_regions.add(r)
            if not all_regions: all_regions = {"77"}
            reg = self.registry_list(sorted(all_regions))
            reg_inns = {r["inn"]: r for r in reg}
            web_inns = {b["inn"]: b for b in all_web if b.get("inn")}

            # джойн: веб → реестр (по ИНН/названию)
            joined_inns = set()
            for b in all_web:
                inn = b.get("inn")
                rr = reg_inns.get(inn) if inn else None
                if not rr:
                    # fallback: name match
                    k = " ".join(sorted(_tok(b.get("site_brand") or b.get("title",""))))
                    for r in reg:
                        if r["inn"] not in joined_inns and k and k == " ".join(sorted(_tok(r["name"]))):
                            rr = r; break
                if rr:
                    b["inn"] = rr["inn"]; b["egrul_name"] = rr["name"]
                    b["is_ip"] = rr["is_ip"]; b["okveds"] = rr["okveds"]
                    if rr["is_ip"]: b["director"] = rr["fio"]
                    joined_inns.add(rr["inn"])
                rows.append(b)
            self.log(f"Реестр: дал ИНН/директора {len(joined_inns)} веб-клиникам (всего в реестре {len(reg)})")
            # ВАЖНО: реестр-приложение (клиники без веб-присутствия) в Excel НЕ добавляем —
            # пользователю нужны только найденные в Яндексе (топ N по городу), а не вся база МСП.
        else:
            rows = all_web

        # ФАЗА 3: DaData только по контактным (у кого ИНН и веб-присутствие)
        self._run_dadata(rows)

        # ФАЗА 4: 2ГИС + РКН — ТОЛЬКО по контактным (с веб-присутствием),
        # а не по всему реестру (12k мёртвых лидов без сайта)
        def _contact(b): return bool(b.get("domain") or b.get("phones") or b.get("emails"))
        for city in self.cfg["cities"]:
            if self.should_stop(): break
            cr = [b for b in rows if b.get("city","") == city and _contact(b)]
            self.log(f"=== {city}: 2ГИС по {len(cr)} контактным ===")
            self.twogis(cr, city)
        if self.cfg.get("rkn"):
            # РКН — только по контактным (с веб-присутствием), а не по всему реестру
            rkn_todo = [b for b in rows if (b.get("domain") or b.get("phones") or b.get("emails")) and b.get("inn")]
            self.log(f"=== РКН: {len(rkn_todo)} контактных (из {len(rows)}) ===")
            try:
                from rkn_enricher import enrich_batch
                enrich_batch(rkn_todo, log_callback=self.log)
            except Exception as e: self.log(f"  РКН: {e}")

        path = self.save_xlsx(rows)
        return path, len(rows)

    # ---------- Excel ----------
    @staticmethod
    def _guess_email(b):
        """Маршрут 4: из найденного email домена + ФИО директора собрать предполагаемый email.
        Только если есть реальный email того же домена (знаем паттерн) и ФИО директора."""
        from transliterate import transliterate_to_russian  # есть обратная? используем простую translit
        emails=b.get("emails",[]); director=b.get("director","")
        domain=b.get("domain","")
        if not emails or not director or not domain: return ""
        # берём email на домене сайта
        same=[e for e in emails if domain.split(".")[0] in e or e.endswith("@"+domain)]
        if not same: return ""
        local=same[0].split("@")[0]
        # транслит фамилии/имени директора (рус → лат)
        TR={'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z','и':'i',
            'й':'y','к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r','с':'s','т':'t',
            'у':'u','ф':'f','х':'kh','ц':'ts','ч':'ch','ш':'sh','щ':'sch','ъ':'','ы':'y',
            'ь':'','э':'e','ю':'yu','я':'ya',' ':''}
        def tr(s): return "".join(TR.get(c,c) for c in s.lower())
        parts=director.split()
        if len(parts)<2: return ""
        fam=tr(parts[0]); im=tr(parts[1])
        if not fam: return ""
        dom_at="@"+(same[0].split("@")[1])
        # угадываем формат по образцу local (если info/admin — берём стандартный f.lastname)
        if local in ("info","admin","mail","clinic","reception","zapis"):
            guess=f"{im[0]}.{fam}{dom_at}" if im else f"{fam}{dom_at}"
        elif "." in local:   # формат name.surname
            guess=f"{im}.{fam}{dom_at}"
        else:
            guess=f"{fam}{dom_at}"
        return guess

    def save_xlsx(self, rows):
        for b in rows: b.setdefault("rank",99999)
        # ── дедуп: одна клиника не должна повторяться (по ИНН, иначе по домену) ──
        seen=set(); uniq=[]
        for b in rows:
            key = b.get("inn") or b.get("domain") or id(b)
            if key in seen: continue
            seen.add(key); uniq.append(b)
        if len(uniq) < len(rows):
            self.log(f"Дедуп: {len(rows)} → {len(uniq)} (убрано {len(rows)-len(uniq)} повторов)")
        rows = uniq
        # ── моб. из мессенджеров (t.me/+7…, wa.me/7…) в отдельное поле ──
        for b in rows:
            mob=""
            for s in b.get("socials", []):
                m=re.search(r"(?:wa\.me/|t\.me/\+?)(\d{10,11})", s)
                if m:
                    dg=m.group(1)
                    if len(dg)==11 and dg[0] in "78": mob="+7"+dg[1:]
                    elif len(dg)==10: mob="+7"+dg
                    if mob: break
            b["msg_phone"]=mob
        # ── маршрут 4: предполагаемый email директора по паттерну ──
        for b in rows:
            b["guess_email"]=self._guess_email(b)
        # группируем по городу (в порядке из конфига), внутри города — по позиции Яндекса
        city_order = {c: i for i, c in enumerate(self.cfg.get("cities", []))}
        def yapos(b):
            try: return int(b.get("yandex_pos") or 99999)
            except (ValueError, TypeError): return 99999
        rows.sort(key=lambda x: (city_order.get(x.get("city",""), 99),
                                 yapos(x),
                                 x.get("rank") or 99999))
        wb=openpyxl.Workbook(); ws=wb.active; ws.title="Бизнесы"
        H=["№","Город","Бренд","Название ЕГРЮЛ","ИНН","ОГРН","Директор/ИП","Учредители","Руководство (сайт)",
           "Юр.адрес","Статус","Телефоны","Моб. из мессенджера","Email","Email директора (предпол.)",
           "Мессенджеры","Рейтинг 2ГИС","Отзывы 2ГИС","Яндекс↑","Сайт","ОКВЭД",
           "Отв. за ПД (РКН)","Телефон РКН","Email РКН"]
        W=[5,16,24,34,14,16,28,30,32,42,12,24,18,26,24,22,11,11,8,24,26,28,16,26]
        hf=Font(bold=True,color="FFFFFF"); hl=PatternFill("solid",start_color="1F4E79")
        for c,(h,w) in enumerate(zip(H,W),1):
            cell=ws.cell(1,c,h); cell.font=hf; cell.fill=hl
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width=w
        for i,b in enumerate(rows,1):
            ws.append([i, b.get("city",""), b.get("site_brand",""), b.get("egrul_name",""),
                       b.get("inn",""), b.get("ogrn",""), b.get("director",""), b.get("founders",""),
                       "; ".join(b.get("leaders",[])),
                       b.get("reg_address",""), b.get("status",""),
                       ", ".join(b.get("phones",[])), b.get("msg_phone",""),
                       ", ".join(b.get("emails",[])[:3]), b.get("guess_email",""),
                       ", ".join(b.get("socials",[])), b.get("twogis_rating",""), b.get("twogis_reviews",""),
                       b.get("yandex_pos",""), ("https://"+b["domain"]) if b.get("domain") else "",
                       ", ".join(b.get("okveds",[])[:6]) if b.get("okveds") else "",
                       b.get("rkn_name",""), b.get("rkn_phone",""), b.get("rkn_email","")])
        ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(H))}{len(rows)+1}"
        from datetime import datetime
        fn=f"БАЗА_{self.cfg['business']}_{'_'.join(self.cfg['cities'])}.xlsx".replace(" ","_")
        wb.save(fn); return fn
